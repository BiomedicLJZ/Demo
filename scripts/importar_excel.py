# /// script
# dependencies = ["openpyxl"]
# ///
"""Genera un seed SQL desde un Excel de Intelisis (hojas Articulos, Existencias, Clientes, Ventas_Detalle).
Uso: uv run scripts/importar_excel.py [libro.xlsx] [salida.sql]      (por defecto 'Data Final.xlsx' -> db/seed.sql)
Luego, Postgres: docker compose exec -T db psql -v ON_ERROR_STOP=1 -U postgres -d ventas -f /sql/seed.sql
      SQLite:    uv run scripts/init_sqlite.py fevasa.db salida.sql
Impuesto = Impuesto1/100 por producto (Impuesto2/3 se ignoran). Ignora sucursales, vendedores y CRM: el esquema no tiene dónde guardarlos.
Un PrecioLista <= 1 es un marcador, no un precio: se sustituye por el precio de venta más frecuente (> 1) de Ventas_Detalle;
sin ventas que lo respalden el producto se carga inactivo. Sin PrecioLista el producto se omite."""
import sys
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
MARCADOR = 1  # precio de lista que equivale a «sin precio»
q = lambda s: "'" + str(s).replace("'", "''") + "'"


def filas(wb, hoja):
    it = wb[hoja].iter_rows(values_only=True)
    cols = next(it)
    return [dict(zip(cols, r)) for r in it if any(c is not None for c in r)]


def numero(d: Decimal) -> str:
    return format(d.normalize(), "f")


libro = Path(sys.argv[1] if len(sys.argv) > 1 else RAIZ / "Data Final.xlsx")
salida = Path(sys.argv[2] if len(sys.argv) > 2 else RAIZ / "db" / "seed.sql")
wb = openpyxl.load_workbook(libro, read_only=True, data_only=True)

stock = defaultdict(Decimal)  # suma de todos los almacenes; una combinación ausente cuenta como 0
for e in filas(wb, "Existencias"):
    assert e["Disponible"] >= 0, e
    stock[e["Articulo"]] += Decimal(str(e["Disponible"])).quantize(Decimal("0.001"))  # quita el ruido de punto flotante

ventas = defaultdict(Counter)  # precio de venta -> veces, por artículo
if "Ventas_Detalle" in wb.sheetnames:
    for d in filas(wb, "Ventas_Detalle"):
        if d["Precio"] and d["Precio"] > MARCADOR:
            ventas[d["Articulo"]][round(d["Precio"], 2)] += 1

prods = []
for a in filas(wb, "Articulos"):
    sku, precio, activo = a["Articulo"], a["PrecioLista"], a["SeVende"] == 1
    if precio is None:
        print(f"OMITIDO {sku} {a['Descripcion']}: sin precio")
        continue
    if precio <= MARCADOR:
        if ventas[sku]:
            precio = max(ventas[sku].items(), key=lambda kv: (kv[1], kv[0]))[0]  # más frecuente; empate: el mayor
            print(f"PRECIO {sku} {a['Descripcion']}: 1 -> {precio} (de {sum(ventas[sku].values())} ventas)")
        else:
            activo = False
            print(f"INACTIVO {sku} {a['Descripcion']}: precio 1 y sin ventas que lo respalden")
    unidad = (a["Unidad"] or "").strip()
    nombre = a["Descripcion"] + (f" [{unidad}]" if unidad and unidad.upper() != "PZA" else "")  # el precio es por esta unidad
    prods.append(f"  ({q(sku)}, {q(nombre)}, {round(precio, 2)}, {numero(stock[sku])}, "
                 f"{'TRUE' if activo else 'FALSE'}, {Decimal(str(a['Impuesto1'])) / 100})")

# sin email en el origen: se sintetiza por ID (los nombres se repiten, p. ej. PÚBLICO EN GENERAL)
clis = [f"  ({q(c['Nombre'])}, {q('cliente-' + str(c['Cliente']) + '@demo.invalid')}, NULL)"
        for c in filas(wb, "Clientes")]

salida.write_text(
    "TRUNCATE comprobantes, pagos, pedido_items, pedidos, productos, clientes RESTART IDENTITY CASCADE;\n"
    "ALTER SEQUENCE comprobante_seq RESTART;\n\n"
    "INSERT INTO productos (sku, nombre, precio, stock, activo, impuesto) VALUES\n" + ",\n".join(prods) + ";\n\n"
    "INSERT INTO clientes (nombre, email, telefono) VALUES\n" + ",\n".join(clis) + ";\n")
print(f"{salida}: {len(prods)} productos, {len(clis)} clientes")
